"""Excel reading and header inference (spec §5).

Deterministic: normalise, match against a synonym table, report what is left over.
No model call — see spec §5 and CLAUDE.md Rule 5.
"""

import datetime as dt
import re
from dataclasses import dataclass, field

from .models import Problem, RoleSpec, UserSpec

# Field -> accepted header spellings, already normalised.
SYNONYMS: dict[str, set[str]] = {
    "username": {"user", "username", "userid", "username", "bname", "sapuser", "login",
                 "logonname", "userlogin", "sapuserid", "account"},
    "last_name": {"lastname", "surname", "familyname", "nachn", "name2", "lname"},
    "first_name": {"firstname", "givenname", "forename", "vorna", "name1", "fname"},
    "email": {"email", "emailaddress", "mail", "smtpaddr", "smtpaddress", "internetaddress"},
    "user_type": {"usertype", "type", "ustyp", "logontype"},
    "user_group": {"usergroup", "group", "class", "usergrp"},
    "valid_from": {"validfrom", "validityfrom", "startdate", "gltgv", "from"},
    "valid_to": {"validto", "validityto", "enddate", "gltgb", "to", "until"},
    "initial_password": {"password", "initialpassword", "initpassword", "pwd"},
    "snc_name": {"sncname", "snc", "sncprintablename", "sncpname"},
    "roles": {"role", "roles", "rolename", "agrname", "singlerole", "compositerole",
              "saprole", "authorisation", "authorization", "profile"},
}

_ROLE_SPLIT = re.compile(r"[,;\n]+")
_SCAN_ROWS = 10  # how far down to look for the header row


@dataclass
class ParseResult:
    users: list[UserSpec] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)  # field -> original header text
    unmapped: list[str] = field(default_factory=list)
    problems: list[Problem] = field(default_factory=list)
    header_row: int = 0


def normalise(text: object) -> str:
    """'User ID' / 'user_id' / 'USERID' all collapse to 'userid'."""
    return re.sub(r"[^a-z0-9]", "", str(text or "").strip().lower())


def match_header(header: object) -> str | None:
    key = normalise(header)
    if not key:
        return None
    for field_name, spellings in SYNONYMS.items():
        if key in spellings:
            return field_name
    return None


def find_header_row(rows: list[list[object]]) -> int:
    """Pick the row with the most synonym hits. Real exports carry title banners."""
    best_row, best_hits = 0, 0
    for index, row in enumerate(rows[:_SCAN_ROWS]):
        hits = sum(1 for cell in row if match_header(cell))
        if hits > best_hits:
            best_row, best_hits = index, hits
    return best_row


def normalise_date(value: object) -> str:
    """Accept what spreadsheets actually contain; emit YYYYMMDD or '' or the raw text."""
    if value is None or value == "":
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y%m%d")
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y%m%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(text, pattern).strftime("%Y%m%d")
        except ValueError:
            continue
    return text  # left as-is so the validator reports it against a real row


def split_roles(cell: object) -> list[str]:
    if cell is None:
        return []
    return [part.strip().upper() for part in _ROLE_SPLIT.split(str(cell)) if part.strip()]


def parse_rows(rows: list[list[object]], header_row: int, mapping: dict[str, int]) -> ParseResult:
    """Group data rows into UserSpec. Handles both sheet layouts from spec §5.

    Layout A (roles in one delimited cell) and Layout B (one row per user-role pair)
    both fall out of grouping by username — a user in three rows becomes one user
    with three roles, never three users.
    """
    result = ParseResult(header_row=header_row)
    by_name: dict[str, UserSpec] = {}

    def cell(row: list[object], name: str) -> str:
        index = mapping.get(name)
        if index is None or index >= len(row):
            return ""
        value = row[index]
        return "" if value is None else str(value).strip()

    for offset, row in enumerate(rows[header_row + 1 :]):
        sheet_row = header_row + offset + 2  # 1-based, and past the header
        if not any(str(c).strip() for c in row if c is not None):
            continue

        username = cell(row, "username").upper()
        if not username:
            result.problems.append(Problem((sheet_row,), "username", "missing"))
            continue

        role_dates = (
            normalise_date(row[mapping["valid_from"]]) if "valid_from" in mapping else "",
            normalise_date(row[mapping["valid_to"]]) if "valid_to" in mapping else "",
        )
        roles = [RoleSpec(name, *role_dates) for name in split_roles(cell(row, "roles"))]

        existing = by_name.get(username)
        if existing is None:
            by_name[username] = UserSpec(
                username=username,
                last_name=cell(row, "last_name"),
                first_name=cell(row, "first_name"),
                email=cell(row, "email"),
                user_type=(cell(row, "user_type") or "A").upper()[:1],
                user_group=cell(row, "user_group").upper(),
                valid_from=role_dates[0],
                valid_to=role_dates[1],
                initial_password=cell(row, "initial_password"),
                snc_name=cell(row, "snc_name"),
                roles=roles,
                source_rows=[sheet_row],
            )
            continue

        existing.source_rows.append(sheet_row)
        known = {r.name for r in existing.roles}
        existing.roles.extend(r for r in roles if r.name not in known)
        _merge_scalars(existing, row, cell, sheet_row, result)

    result.users = list(by_name.values())
    return result


def _merge_scalars(existing, row, cell, sheet_row, result) -> None:
    """Later rows may fill blanks, but must not silently contradict earlier ones."""
    for name in ("last_name", "first_name", "email", "user_group", "initial_password", "snc_name"):
        incoming = cell(row, name)
        if not incoming:
            continue
        current = getattr(existing, name)
        if not current:
            setattr(existing, name, incoming)
        elif current != incoming:
            result.problems.append(
                Problem(
                    tuple(existing.source_rows),
                    name,
                    f"'{existing.username}' has conflicting values '{current}' and '{incoming}'",
                )
            )


def parse_file(path: str, sheet: str | None = None) -> ParseResult:
    """Dispatch on extension. Both formats reach the same header inference."""
    if path.lower().endswith(".csv"):
        return parse_csv(path)
    return parse_workbook(path, sheet)


def parse_csv(path: str) -> ParseResult:
    import csv

    # utf-8-sig: Excel writes a BOM, which would otherwise corrupt the first header.
    with open(path, newline="", encoding="utf-8-sig") as handle:
        rows = [list(row) for row in csv.reader(handle)]
    return _from_rows(rows)


def parse_workbook(path: str, sheet: str | None = None) -> ParseResult:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True, read_only=True)
    worksheet = book[sheet] if sheet else book.worksheets[0]
    rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
    book.close()
    return _from_rows(rows)


def _from_rows(rows: list[list[object]]) -> ParseResult:
    if not rows:
        return ParseResult(problems=[Problem((), "sheet", "file is empty")])

    header_row = find_header_row(rows)
    headers = rows[header_row]

    mapping: dict[str, int] = {}
    original: dict[str, str] = {}
    unmapped: list[str] = []
    for index, header in enumerate(headers):
        matched = match_header(header)
        if matched and matched not in mapping:
            mapping[matched] = index
            original[matched] = str(header).strip()
        elif str(header or "").strip():
            unmapped.append(str(header).strip())

    result = parse_rows(rows, header_row, mapping)
    result.mapping = original
    result.unmapped = unmapped
    result.header_row = header_row

    for required in ("username", "roles"):
        if required not in mapping:
            result.problems.append(
                Problem((header_row + 1,), required, "no column matched this field")
            )
    return result
